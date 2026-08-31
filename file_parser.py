import io


def _decode_bytes(data: bytes) -> str:
    """字节流解码，自动适配 UTF-8 / GB18030 编码。"""
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def parse_pdf(data: bytes) -> str:
    """解析 PDF 为纯文本。"""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n\n".join(parts)


def parse_docx(data: bytes) -> str:
    """解析 .docx 为纯文本（含段落与表格）。"""
    try:
        import docx
    except ImportError as exc:  # noqa: BLE001
        raise RuntimeError("环境未安装 python-docx，无法解析 .docx 文件") from exc

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(cells).strip()
            if line.strip(" |"):
                parts.append(line)
    return "\n".join(parts)


def parse_xlsx(data: bytes) -> str:
    """解析 .xlsx 为纯文本（每个 sheet 逐行拼接）。"""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"【Sheet：{ws.title}】")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            line = " | ".join(cells).strip()
            if line.strip(" |"):
                parts.append(line)
    return "\n".join(parts)


def parse_xls(data: bytes) -> str:
    """解析 .xls 为纯文本。"""
    import xlrd

    wb = xlrd.open_workbook(file_contents=data)
    parts = []
    for sh in wb.sheets():
        parts.append(f"【Sheet：{sh.name}】")
        for r in range(sh.nrows):
            cells = [str(sh.cell_value(r, c)) for c in range(sh.ncols)]
            line = " | ".join(cells).strip()
            if line.strip(" |"):
                parts.append(line)
    return "\n".join(parts)


def parse_uploaded_file(filename: str, data: bytes) -> str:
    """根据扩展名解析上传文件为纯文本。"""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return parse_pdf(data)
    if name.endswith(".docx"):
        return parse_docx(data)
    if name.endswith(".xlsx"):
        return parse_xlsx(data)
    if name.endswith(".xls"):
        return parse_xls(data)
    # 默认按文本处理
    return _decode_bytes(data)
